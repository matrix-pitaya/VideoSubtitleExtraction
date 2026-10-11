"""Local, resumable batch transcription and delivery exports."""
from __future__ import annotations

import json
import re
from pathlib import Path
from datetime import datetime
from uuid import uuid4
from zipfile import ZipFile, ZIP_DEFLATED

ROOT = Path(__file__).resolve().parent / "data"
MEDIA_EXTENSIONS = {".mp4", ".mov", ".mkv", ".avi", ".webm", ".m4v", ".mp3", ".wav", ".m4a", ".aac", ".flac"}
FINISHED_STATUSES = {"完成", "无口播"}


def transcript_paragraphs(segments: list[dict]) -> str:
    """Group subtitle fragments into readable paragraphs without rewriting words."""
    paragraphs, current = [], ""
    needs_punctuation = len(segments) > 1 and not any(re.search(r"[，。！？,!?;；]", s["text"]) for s in segments)

    def finish(value):
        return value.rstrip("，,") + "。" if needs_punctuation and value and not re.search(r"[。！？.!?]$", value) else value
    previous_end = None
    for segment in segments:
        text = segment["text"].strip()
        if not text:
            continue
        if current and previous_end is not None and segment["start"] - previous_end >= 2:
            paragraphs.append(finish(current))
            current = ""
        separator = " " if current and current[-1].isascii() and current[-1].isalnum() and text[0].isascii() and text[0].isalnum() else ""
        current += separator + text
        if needs_punctuation and re.search(r"[\u3400-\u9fff]", text) and not re.search(r"[，。！？,:;.!?]$", current):
            current += "，"
        previous_end = segment["end"]
        if len(current) >= 150 and (needs_punctuation or re.search(r"[。！？.!?][\"'”’）)]?$", current)):
            paragraphs.append(finish(current))
            current = ""
        elif len(current) >= 220:
            paragraphs.append(finish(current))
            current = ""
    if current:
        paragraphs.append(finish(current))
    return "\n\n".join(paragraphs)


def safe_name(value: str) -> str:
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", value).strip(" .")[:80] or "未命名"


def extract_links(text: str) -> list[str]:
    return list(dict.fromkeys(s.rstrip(".,，。；;！!）)]】") for s in re.findall(r"https?://[^\s<>\"]+", text)))


def save_order(order: dict) -> None:
    path = ROOT / order["id"] / "order.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".tmp")
    temporary.write_text(json.dumps(order, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def create_order(name: str) -> dict:
    order = {"id": datetime.now().strftime("%Y%m%d-%H%M%S-") + uuid4().hex[:8], "name": name or "新订单", "items": []}
    save_order(order)
    return order


def load_orders() -> list[dict]:
    orders = []
    for path in sorted(ROOT.glob("*/order.json"), reverse=True):
        try:
            orders.append(json.loads(path.read_text(encoding="utf-8")))
        except (OSError, ValueError):
            continue
    return orders


def add_item(order: dict, source: str, title: str, local_path: str = "") -> dict:
    item = {"id": uuid4().hex[:12], "source": source, "title": title, "path": local_path,
            "status": "待处理", "error": "", "segments": [], "text": ""}
    order["items"].append(item)
    save_order(order)
    return item


def download(item: dict, folder: Path, cookie_path: str = "") -> str:
    from link_download import is_douyin, download_douyin, is_xiaohongshu, download_xiaohongshu
    from link_download import is_supported_link
    if not is_supported_link(item["source"]):
        raise ValueError("链接提取目前仅支持抖音、B站、小红书，请使用这三个平台的分享链接或导入本地视频")
    if is_xiaohongshu(item["source"]):
        return download_xiaohongshu(item, folder, cookie_path)
    platform_error = ""
    if is_douyin(item["source"]):
        try:
            return download_douyin(item, folder)
        except Exception as exc:
            platform_error = str(exc)
    from yt_dlp import YoutubeDL
    folder.mkdir(parents=True, exist_ok=True)
    options = {"outtmpl": str(folder / "%(id)s.%(ext)s"), "format": "bestaudio/best",
               "noplaylist": True, "socket_timeout": 30, "retries": 2,
               "quiet": True, "windowsfilenames": True}
    if cookie_path:
        options["cookiefile"] = cookie_path
    try:
        with YoutubeDL(options) as downloader:
            info = downloader.extract_info(item["source"], download=True)
            if not info or info.get("entries") is not None:
                raise ValueError("请提供单个视频链接，不能使用合集或列表链接")
            item["title"] = info.get("title") or item["title"]
            path = downloader.prepare_filename(info)
    except Exception as exc:
        if platform_error:
            raise ValueError(f"下载失败。抖音分享解析：{platform_error}；备用下载器：{exc}") from exc
        raise
    if not Path(path).is_file():
        raise FileNotFoundError("未获取到媒体文件，请保存视频后本地导入")
    return path


def process_item(order: dict, item: dict, model, language: str | None, cookie_path: str = "", on_progress=None, mode: str = "audio") -> None:
    is_link = item["source"].startswith(("https://", "http://"))
    download_folder = ROOT / order["id"] / "media" / item["id"]
    try:
        item.update(status="处理中", error="", cleanup_error="")
        save_order(order)
        if not item["path"] or not Path(item["path"]).is_file():
            if not is_link:
                raise FileNotFoundError("本地视频不存在，请重新导入或恢复原文件")
            item["status"] = "解析并下载链接"
            save_order(order)
            if on_progress:
                on_progress("解析并下载链接")
            item["path"] = download(item, download_folder, cookie_path)
            save_order(order)
        item["status"] = "提取口播字幕"
        save_order(order)
        if on_progress:
            on_progress("提取口播字幕")
        item["segments"] = []
        item["text_origin"] = "口播转写"
        if mode != "ocr":
            try:
                segments, info = model.transcribe(item["path"], language=language, beam_size=5, vad_filter=True,
                                                  initial_prompt="请用简体中文转写口播，保留自然标点。" if language == "zh" else None)
                item["segments"] = [{"start": s.start, "end": s.end, "text": s.text.strip()} for s in segments if s.text.strip()]
                item["language"] = info.language
            except ValueError as exc:
                if "audio" not in str(exc).lower():
                    raise
        if mode == "ocr":
            from screen_text import extract_screen_text
            item["status"] = "识别画面文字"
            save_order(order)
            if on_progress:
                on_progress("无口播，正在识别画面文字" if mode == "auto" else "正在识别画面文字")
            item["segments"] = extract_screen_text(item["path"], on_progress)
            item["text_origin"] = "画面文字 OCR（1秒采样，时间轴近似）"
        item["text"] = transcript_paragraphs(item["segments"])
        if not item["segments"]:
            item["status"] = "无口播" if mode != "ocr" else "失败"
            item["note"] = "未检测到口播，文案留空" if mode != "ocr" else "未检测到清晰画面文字"
            if mode == "ocr":
                item["error"] = item["note"]
        else:
            item["status"] = "完成"
            item["note"] = ""
    except Exception as exc:
        item.update(status="失败", error=str(exc))
    finally:
        if is_link:
            try:
                cleanup_download(download_folder)
                item["path"] = ""
            except OSError as exc:
                item["cleanup_error"] = f"临时媒体删除失败：{exc}"
        save_order(order)


def cleanup_download(folder: Path) -> None:
    """Delete only files inside this item's owned media directory, including partial downloads."""
    base = ROOT.resolve()
    resolved = folder.resolve()
    if not resolved.is_relative_to(base) or len(resolved.relative_to(base).parts) != 3 or resolved.relative_to(base).parts[1] != "media":
        raise OSError("拒绝清理订单媒体目录之外的路径")
    if folder.is_symlink():
        raise OSError("拒绝清理符号链接目录")
    if not folder.exists():
        return
    paths = list(folder.rglob("*"))
    # Validate every target before removing any file.
    if any(p.is_symlink() or not p.resolve().is_relative_to(resolved) for p in paths):
        raise OSError("临时目录包含外部链接，未执行清理")
    for path in sorted(paths, key=lambda p: len(p.parts), reverse=True):
        if path.is_dir():
            path.rmdir()
        else:
            path.unlink()
    folder.rmdir()


def timestamp(seconds: float) -> str:
    milliseconds = max(0, round(seconds * 1000))
    hours, rest = divmod(milliseconds, 3600000)
    minutes, rest = divmod(rest, 60000)
    seconds, milliseconds = divmod(rest, 1000)
    return f"{hours:02}:{minutes:02}:{seconds:02},{milliseconds:03}"


def srt(segments: list[dict]) -> str:
    return "\n\n".join(f"{i}\n{timestamp(s['start'])} --> {timestamp(s['end'])}\n{s['text']}" for i, s in enumerate(segments, 1)) + "\n"


def export_order(order: dict) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    folder = ROOT / order["id"] / "exports" / "delivery"
    folder.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "文案汇总"
    sheet.append(["编号", "标题", "文案"])
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "C2"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"
    sheet.column_dimensions["A"].width = 8
    sheet.column_dimensions["B"].width = 42
    sheet.column_dimensions["C"].width = 125
    sheet.row_dimensions[1].height = 28
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    for style_name in ("Normal", "Title", "Heading 1"):
        style = document.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal = document.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.first_line_indent = Pt(22)
    normal.paragraph_format.space_after = Pt(6)
    heading = document.styles["Heading 1"]
    heading.font.size = Pt(13)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor.from_string("245A81")
    heading.paragraph_format.first_line_indent = Pt(0)
    heading.paragraph_format.space_before = Pt(16)
    heading.paragraph_format.space_after = Pt(8)
    heading.paragraph_format.keep_with_next = True
    document.styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    document.styles["Title"].font.size = Pt(20)
    for style in document.styles:
        if style.type == 1:
            ppr = style._element.find(qn("w:pPr"))
            if ppr is not None:
                for border in list(ppr.findall(qn("w:pBdr"))):
                    ppr.remove(border)
    title = document.add_paragraph(order["name"] + " 文案合集", "Title")
    title.paragraph_format.first_line_indent = Pt(0)
    title.paragraph_format.space_after = Pt(18)
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.comments = ""
    footer = section.footer.paragraphs[0]
    footer.paragraph_format.first_line_indent = Pt(0)
    footer.alignment = 2
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    files = []
    for index, item in enumerate(order["items"], 1):
        # Explicit string cell types prevent spreadsheet formula execution.
        text = item["text"] if item["status"] == "完成" else ""
        # Customer deliverables include completed items and confirmed empty speech only.
        if item["status"] not in FINISHED_STATUSES:
            continue
        if len(text) > 32767:
            raise ValueError(f"视频 {index} 的文案超过 Excel 单元格 32767 字符限制，请分拆视频或订单后导出")
        sheet.append([index, item["title"], text.replace("\n\n", "\n")])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
        lines = max(len(item["title"]) // 24 + 1, sum(max(1, (len(line) + 57) // 58) for line in text.splitlines()))
        sheet.row_dimensions[sheet.max_row].height = min(409, max(54, lines * 16 + 12))
        document.add_heading(f"{index}. {item['title']}", 1)
        for paragraph in re.split(r"\n\s*\n", text):
            if paragraph.strip():
                document.add_paragraph(paragraph.strip().replace("\n", ""))
        if not text:
            continue
        stem = f"{index:03d}_{safe_name(item['title'])}"
        subtitle = folder / (stem + ".srt")
        subtitle.write_text(srt(item["segments"]), encoding="utf-8-sig")
        transcript = folder / (stem + ".txt")
        transcript.write_text(text, encoding="utf-8-sig")
        files.extend([subtitle, transcript])
    for ws in workbook:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="245A81")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.font = Font(name="Microsoft YaHei", size=10, color="202020")
                cell.border = Border(bottom=Side(style="hair", color="D7E0E7"))
                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F3F7FA")
        ws.print_area = ws.dimensions
    excel = folder / "文案汇总.xlsx"
    word = folder / "文案汇总.docx"
    workbook.save(excel)
    document.save(word)
    files.extend([excel, word])
    archive = folder / (safe_name(order["name"]) + "_交付包.zip")
    with ZipFile(archive, "w", ZIP_DEFLATED) as zf:
        for path in files:
            zf.write(path, path.name)
    return archive
