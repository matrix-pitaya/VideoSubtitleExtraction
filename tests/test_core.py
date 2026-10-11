import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from types import SimpleNamespace
from zipfile import ZipFile
import core


class WorkflowTests(unittest.TestCase):
    def test_paragraph_formatting_retains_words_and_adds_breaks(self):
        import re
        segments = [{"start": n, "end": n + 1, "text": "这是口播正文" * 5} for n in range(10)]
        result = core.transcript_paragraphs(segments)
        self.assertIn("\n\n", result)
        self.assertEqual(re.sub(r"[，。\s]", "", result), "".join(s["text"] for s in segments))

    def test_link_media_cleanup_on_success_and_failure(self):
        class Model:
            def transcribe(self, path, **kwargs):
                return iter([SimpleNamespace(start=0, end=1, text="口播")]), SimpleNamespace(language="zh")
        for fail in [False, True]:
            with self.subTest(fail=fail), tempfile.TemporaryDirectory() as tmp, patch.object(core, "ROOT", Path(tmp)):
                order = core.create_order("链接订单")
                item = core.add_item(order, "https://example.com/video", "视频")
                def downloader(item, folder, cookie, channels_parser="", channels_key=""):
                    folder.mkdir(parents=True)
                    path = folder / "video.mp4"
                    path.write_bytes(b"test")
                    (folder / "video.part").touch()
                    if fail:
                        raise ValueError("下载中断")
                    return str(path)
                with patch.object(core, "download", side_effect=downloader):
                    core.process_item(order, item, Model(), "zh")
                self.assertFalse((Path(tmp) / order["id"] / "media" / item["id"]).exists())
                self.assertEqual(item["path"], "")
                self.assertEqual(item["status"], "失败" if fail else "完成")
                if not fail:
                    self.assertEqual(item["text"], "口播")

    def test_cleanup_rejects_external_directory(self):
        with tempfile.TemporaryDirectory() as tmp, patch.object(core, "ROOT", Path(tmp) / "data"):
            original = Path(tmp) / "original"
            original.mkdir()
            video = original / "video.mp4"
            video.touch()
            with self.assertRaises(OSError):
                core.cleanup_download(original)
            self.assertTrue(video.exists())

    def test_links_and_timestamps(self):
        self.assertEqual(core.extract_links("分享 https://v.douyin.com/abc/。\nhttps://v.douyin.com/abc/"), ["https://v.douyin.com/abc/"])
        self.assertEqual(core.timestamp(59.9996), "00:01:00,000")
        self.assertIn("00:00:01,250 --> 00:00:02,500", core.srt([{"start": 1.25, "end": 2.5, "text": "你好"}]))

    def test_failure_and_resume(self):
        with tempfile.TemporaryDirectory() as tmp, patch.object(core, "ROOT", Path(tmp)):
            order = core.create_order("测试")
            item = core.add_item(order, "https://example.com/a", "视频")
            with patch.object(core, "download", side_effect=ValueError("无法下载")):
                core.process_item(order, item, None, "zh")
            self.assertEqual(core.load_orders()[0]["items"][0]["status"], "失败")
            self.assertEqual(item["error"], "无法下载")

    def test_delivery_content(self):
        from openpyxl import load_workbook
        from docx import Document
        with tempfile.TemporaryDirectory() as tmp, patch.object(core, "ROOT", Path(tmp)):
            order = core.create_order("客户")
            item = core.add_item(order, "local.mp4", "=1+1")
            item.update(status="完成", text="你好世界", segments=[{"start": 0, "end": 1.5, "text": "你好世界"}])
            archive = core.export_order(order)
            with ZipFile(archive) as zf:
                self.assertEqual(len(zf.namelist()), 4)
                self.assertIn("你好世界", zf.read("001_=1+1.srt").decode("utf-8-sig"))
            folder = archive.parent
            book = load_workbook(folder / "文案汇总.xlsx")
            self.assertEqual(book.active["B2"].data_type, "s")
            self.assertEqual(book.active["C2"].value, "你好世界")
            self.assertEqual([c.value for c in book.active[1]], ["编号", "标题", "文案"])
            self.assertEqual(book.sheetnames, ["文案汇总"])
            self.assertIn("你好世界", [p.text for p in Document(folder / "文案汇总.docx").paragraphs])

    def test_successful_transcription_is_persisted(self):
        class FakeModel:
            def transcribe(self, path, **kwargs):
                return iter([SimpleNamespace(start=0.2, end=1.2, text=" 测试口播 ")]), SimpleNamespace(language="zh")
        with tempfile.TemporaryDirectory() as tmp, patch.object(core, "ROOT", Path(tmp)):
            media = Path(tmp) / "test.mp4"
            media.touch()
            order = core.create_order("测试")
            item = core.add_item(order, str(media), "测试视频", str(media))
            core.process_item(order, item, FakeModel(), "zh")
            restored = core.load_orders()[0]["items"][0]
            self.assertEqual(restored["status"], "完成")
            self.assertEqual(restored["text"], "测试口播")
            self.assertEqual(restored["segments"][0]["start"], 0.2)
            self.assertTrue(media.exists())

    def test_customer_export_has_empty_bgm_and_no_internal_data(self):
        from openpyxl import load_workbook
        from docx import Document
        with tempfile.TemporaryDirectory() as tmp, patch.object(core, "ROOT", Path(tmp)):
            order = core.create_order("客户交付")
            bgm = core.add_item(order, "https://private-source.example/video", "BGM视频")
            bgm.update(status="无口播", text="不应导出的旧内容", segments=[])
            archive = core.export_order(order)
            book = load_workbook(archive.parent / "文案汇总.xlsx")
            self.assertIsNone(book.active["C2"].value)
            self.assertEqual(book.active.max_column, 3)
            with ZipFile(archive) as zf:
                self.assertEqual(set(zf.namelist()), {"文案汇总.xlsx", "文案汇总.docx"})
            paragraphs = [p.text for p in Document(archive.parent / "文案汇总.docx").paragraphs]
            self.assertFalse(any("private-source" in p or "不应导出" in p for p in paragraphs))
            self.assertIn("1. BGM视频", paragraphs)


if __name__ == "__main__":
    unittest.main()
