"""Local, resumable batch transcription and delivery exports."""
from __future__ import annotations

import json
import re
from pathlib import Path
from datetime import datetime
from uuid import uuid4
from zipfile import ZipFile, ZIP_DEFLATED

ROOT = Path(__file__).resolve().parent / "data"
MEDIA_EXTENSIONS = {".mp4", ".mov", ".mkv", ".avi", ".webm", ".m4v", ".mp3", ".wav", ".m4a", ".aac", ".flac"}


def safe_name(value: str) -> str:
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", value).strip(" .")[:80] or "未命名"


def extract_links(text: str) -> list[str]:
    return list(dict.fromkeys(s.rstrip(".,，。；;！!）)]】") for s in re.findall(r"https?://[^\s<>\"]+", text)))


def save_order(order: dict) -> None:
    path = ROOT / order["id"] / "order.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".tmp")
    temporary.write_text(json.dumps(order, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def create_order(name: str) -> dict:
    order = {"id": datetime.now().strftime("%Y%m%d-%H%M%S-") + uuid4().hex[:8], "name": name or "新订单", "items": []}
    save_order(order)
    return order


def load_orders() -> list[dict]:
    orders = []
    for path in sorted(ROOT.glob("*/order.json"), reverse=True):
        try:
            orders.append(json.loads(path.read_text(encoding="utf-8")))
        except (OSError, ValueError):
            continue
    return orders


def add_item(order: dict, source: str, title: str, local_path: str = "") -> dict:
    item = {"id": uuid4().hex[:12], "source": source, "title": title, "path": local_path,
            "status": "待处理", "error": "", "segments": [], "text": ""}
    order["items"].append(item)
    save_order(order)
    return item


def download(item: dict, folder: Path, cookie_path: str = "") -> str:
    from link_download import is_douyin, download_douyin
    platform_error = ""
    if is_douyin(item["source"]):
        try:
            return download_douyin(item, folder)
        except Exception as exc:
            platform_error = str(exc)
    from yt_dlp import YoutubeDL
    folder.mkdir(parents=True, exist_ok=True)
    options = {"outtmpl": str(folder / "%(id)s.%(ext)s"), "format": "bestaudio/best",
               "noplaylist": True, "socket_timeout": 30, "retries": 2,
               "quiet": True, "windowsfilenames": True}
    if cookie_path:
        options["cookiefile"] = cookie_path
    try:
        with YoutubeDL(options) as downloader:
            info = downloader.extract_info(item["source"], download=True)
            if not info or info.get("entries") is not None:
                raise ValueError("请提供单个视频链接，不能使用合集或列表链接")
            item["title"] = info.get("title") or item["title"]
            path = downloader.prepare_filename(info)
    except Exception as exc:
        if platform_error:
            raise ValueError(f"下载失败。抖音分享解析：{platform_error}；备用下载器：{exc}") from exc
        raise
    if not Path(path).is_file():
        raise FileNotFoundError("未获取到媒体文件，请保存视频后本地导入")
    return path


def process_item(order: dict, item: dict, model, language: str | None, cookie_path: str = "", on_progress=None, mode: str = "auto") -> None:
    is_link = item["source"].startswith(("https://", "http://"))
    download_folder = ROOT / order["id"] / "media" / item["id"]
    try:
        item.update(status="处理中", error="", cleanup_error="")
        save_order(order)
        if not item["path"] or not Path(item["path"]).is_file():
            if not is_link:
                raise FileNotFoundError("本地视频不存在，请重新导入或恢复原文件")
            item["status"] = "解析并下载链接"
            save_order(order)
            if on_progress:
                on_progress("解析并下载链接")
            item["path"] = download(item, download_folder, cookie_path)
            save_order(order)
        item["status"] = "提取口播字幕"
        save_order(order)
        if on_progress:
            on_progress("提取口播字幕")
        item["segments"] = []
        item["text_origin"] = "口播转写"
        if mode != "ocr":
            try:
                segments, info = model.transcribe(item["path"], language=language, beam_size=5, vad_filter=True)
                item["segments"] = [{"start": s.start, "end": s.end, "text": s.text.strip()} for s in segments if s.text.strip()]
                item["language"] = info.language
            except ValueError as exc:
                if mode != "auto" or "audio" not in str(exc).lower():
                    raise
        if mode == "ocr" or (mode == "auto" and not item["segments"]):
            from screen_text import extract_screen_text
            item["status"] = "识别画面文字"
            save_order(order)
            if on_progress:
                on_progress("无口播，正在识别画面文字" if mode == "auto" else "正在识别画面文字")
            item["segments"] = extract_screen_text(item["path"], on_progress)
            item["text_origin"] = "画面文字 OCR（1秒采样，时间轴近似）"
        item["text"] = "\n".join(s["text"] for s in item["segments"])
        if not item["segments"]:
            raise ValueError("未识别到可提取内容，请检查视频是否有口播或清晰的画面文字")
        item["status"] = "完成"
    except Exception as exc:
        item.update(status="失败", error=str(exc))
    finally:
        if is_link:
            try:
                cleanup_download(download_folder)
                item["path"] = ""
            except OSError as exc:
                item["cleanup_error"] = f"临时媒体删除失败：{exc}"
        save_order(order)


def cleanup_download(folder: Path) -> None:
    """Delete only files inside this item's owned media directory, including partial downloads."""
    base = ROOT.resolve()
    resolved = folder.resolve()
    if not resolved.is_relative_to(base) or len(resolved.relative_to(base).parts) != 3 or resolved.relative_to(base).parts[1] != "media":
        raise OSError("拒绝清理订单媒体目录之外的路径")
    if folder.is_symlink():
        raise OSError("拒绝清理符号链接目录")
    if not folder.exists():
        return
    paths = list(folder.rglob("*"))
    # Validate every target before removing any file.
    if any(p.is_symlink() or not p.resolve().is_relative_to(resolved) for p in paths):
        raise OSError("临时目录包含外部链接，未执行清理")
    for path in sorted(paths, key=lambda p: len(p.parts), reverse=True):
        if path.is_dir():
            path.rmdir()
        else:
            path.unlink()
    folder.rmdir()


def timestamp(seconds: float) -> str:
    milliseconds = max(0, round(seconds * 1000))
    hours, rest = divmod(milliseconds, 3600000)
    minutes, rest = divmod(rest, 60000)
    seconds, milliseconds = divmod(rest, 1000)
    return f"{hours:02}:{minutes:02}:{seconds:02},{milliseconds:03}"


def srt(segments: list[dict]) -> str:
    return "\n\n".join(f"{i}\n{timestamp(s['start'])} --> {timestamp(s['end'])}\n{s['text']}" for i, s in enumerate(segments, 1)) + "\n"


def export_order(order: dict) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from docx import Document
    from docx.oxml.ns import qn
    folder = ROOT / order["id"] / "exports"
    folder.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "文案汇总"
    sheet.append(["编号", "标题", "来源", "状态", "文案", "失败原因", "提取方式"])
    timeline = workbook.create_sheet("时间轴")
    timeline.append(["视频编号", "标题", "开始秒", "结束秒", "字幕"])
    document = Document()
    document.styles["Normal"].font.name = "Microsoft YaHei"
    document.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    document.add_heading(order["name"], 0)
    document.add_paragraph("视频口播转写交付 · 请校对人名、数字和专业术语")
    files = []
    for index, item in enumerate(order["items"], 1):
        # Explicit string cell types prevent spreadsheet formula execution.
        sheet.append([index, item["title"], item["source"], item["status"], item["text"], item["error"], item.get("text_origin", "口播转写")])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
        if item["status"] != "完成":
            continue
        document.add_heading(f"{index}. {item['title']}", 1)
        document.add_paragraph("来源：" + item["source"])
        document.add_paragraph("提取方式：" + item.get("text_origin", "口播转写"))
        for line in item["text"].splitlines():
            document.add_paragraph(line)
        stem = f"{index:03d}_{safe_name(item['title'])}"
        subtitle = folder / (stem + ".srt")
        subtitle.write_text(srt(item["segments"]), encoding="utf-8-sig")
        transcript = folder / (stem + ".txt")
        transcript.write_text(item["text"], encoding="utf-8-sig")
        files.extend([subtitle, transcript])
        for segment in item["segments"]:
            timeline.append([index, item["title"], segment["start"], segment["end"], segment["text"]])
            for cell in timeline[timeline.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    for ws in workbook:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="245A81")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for column, width in zip("ABCDEFG", [12, 35, 45, 18, 85, 45, 35]):
            ws.column_dimensions[column].width = width
    excel = folder / "文案汇总.xlsx"
    word = folder / "文案汇总.docx"
    workbook.save(excel)
    document.save(word)
    files.extend([excel, word])
    manifest = folder / "处理记录.json"
    manifest.write_text(json.dumps(order, ensure_ascii=False, indent=2), encoding="utf-8")
    files.append(manifest)
    archive = folder / (safe_name(order["name"]) + "_交付包.zip")
    with ZipFile(archive, "w", ZIP_DEFLATED) as zf:
        for path in files:
            zf.write(path, path.name)
    return archive
